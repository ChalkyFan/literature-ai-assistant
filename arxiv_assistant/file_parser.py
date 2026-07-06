"""File parser: extract text from uploaded files (PDF, DOCX, PPTX)"""
import os
import logging

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx"}


def extract_text(file_path: str) -> str:
    """Extract text from a file based on its extension. Returns empty string on failure."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        logger.warning("Unsupported file type: {0}".format(ext))
        return ""

    if ext == ".pdf":
        return _extract_pdf_text(file_path)
    elif ext == ".docx":
        return _extract_docx_text(file_path)
    elif ext == ".pptx":
        return _extract_pptx_text(file_path)
    return ""


def _extract_pdf_text(file_path: str) -> str:
    try:
        import fitz
        doc = fitz.open(file_path)
        text = ""
        for page in doc:
            text += page.get_text() + "\n"
        doc.close()
        logger.info("Extracted {0} chars from PDF: {1}".format(len(text), file_path))
        return text
    except Exception as e:
        logger.error("PDF text extraction failed: {0}".format(e))
        return ""


def _extract_docx_text(file_path: str) -> str:
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)
        logger.info("Extracted {0} chars from DOCX: {1}".format(len(text), file_path))
        return text
    except Exception as e:
        logger.error("DOCX text extraction failed: {0}".format(e))
        return ""


def _extract_pptx_text(file_path: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text)
        text = "\n".join(texts)
        logger.info("Extracted {0} chars from PPTX: {1}".format(len(text), file_path))
        return text
    except Exception as e:
        logger.error("PPTX text extraction failed: {0}".format(e))
        return ""


def extract_authors(text: str) -> str:
    """Try to extract author names from document text. Returns comma-separated string."""
    import re
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    # Skip the title (usually first non-empty line)
    # Look for patterns like "Authors:" or "By " or name patterns on early lines
    authors = []
    for i, line in enumerate(lines[:30]):
        # Pattern: "Author: Name1, Name2" or "Authors: Name1, Name2"
        m = re.search(r"(?:Author|Authors|By)[:\s]+(.+)", line, re.IGNORECASE)
        if m:
            candidates = re.split(r"[,;]", m.group(1))
            authors = [c.strip() for c in candidates if c.strip() and len(c.strip()) > 3]
            if authors:
                break
        # Pattern: multiple names with superscript markers like "John Smith1, Jane Doe2"
        if i >= 1 and i <= 5:
            # Check if line looks like author list (contains mostly names and numbers/symbols)
            name_count = len(re.findall(r"[A-Z][a-z]+(?:\s+[A-Z][a-z]+)+", line))
            if name_count >= 2 and len(line) < 300:
                authors = [n.strip() for n in re.split(r"[,;]", line) if n.strip()]
                if all(len(n.split()) >= 2 for n in authors[:3]):
                    break
                else:
                    authors = []
    
    if authors:
        return ", ".join(authors[:20])
    return ""
